import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from schedule import read_schedule, report


class ScheduleTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name) / 'test.docx'

    def make_document(self, title='Расписание на 11 сентября'):
        document = Document()
        document.add_paragraph(title)
        table = document.add_table(rows=0, cols=6)
        for values in [
            ['', '7 д 3-20', '', '8з 3-38', '8гум 2-38а', '8 е 2-17'],
            ['1-7', 'Информатика', '1', 'ин ф(РАМ)/ англ', 'Информатика', 'Алгебра'],
            ['2-8', 'История', '2', 'Анг л /инф(ЕЛЮ)', 'Информатика', 'Информатика'],
        ]:
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
        document.save(self.path)

    def test_classes_subgroups_and_shifts(self):
        self.make_document()
        lessons, present = read_schedule(self.path)
        self.assertEqual([(x.school_class, x.number) for x in lessons],
                         [('7Д', '1-7'), ('8З', '1'), ('8З', '2'), ('8Е', '2')])
        self.assertNotIn('8Г', present)

    def test_wrong_date(self):
        self.make_document()
        answer = report(self.path, date(2026, 9, 10), {})
        self.assertIn('7Д: урок 1-7', answer)

    def test_missing_date(self):
        self.make_document('Расписание')
        self.assertIn('7Д: урок 1-7', report(self.path, date(2026, 9, 11), {}))

    def test_bells_and_absent_class(self):
        self.make_document()
        answer = report(self.path, date(2026, 9, 11), {'1-7': '14:00–14:40'})
        self.assertIn('7Д: урок 1-7, 14:00–14:40', answer)
        self.assertNotIn('7А:', answer)
        self.assertIn('время не настроено', answer)

    def test_chronological_order_across_shifts(self):
        self.make_document()
        answer = report(self.path, date(2026, 9, 11), {})
        entries = [line.split(',')[0] for line in answer.splitlines() if ': урок ' in line]
        self.assertEqual(entries, ['8З: урок 1', '8З: урок 2', '8Е: урок 2', '7Д: урок 1-7'])

    def test_no_informatics(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 1).text = '7А'
        table.cell(1, 0).text = '1'
        table.cell(1, 1).text = 'Алгебра'
        document.save(self.path)
        answer = report(self.path, date(2026, 9, 11), {})
        self.assertIn('нет уроков информатики', answer)
        self.assertNotIn('7А:', answer)

    def test_inform_abbreviation_with_group(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 1).text = '7А'
        table.cell(1, 0).text = '1'
        table.cell(1, 1).text = 'Информ(гр. ВМИ чз)'
        table.cell(2, 0).text = '2'
        table.cell(2, 1).text = 'Информация о занятиях'
        document.save(self.path)
        answer = report(self.path, date(2026, 9, 11), {'1': '08:20–09:00'})
        self.assertIn('7А: урок 1, 08:20–09:00 — Информ(гр. ВМИ чз)', answer)
        self.assertNotIn('урок 2', answer)



if __name__ == '__main__':
    unittest.main()
